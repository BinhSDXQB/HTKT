#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script tạo câu SQL INSERT để import dữ liệu "Ngày vào công chức, cơ quan NN" từ Excel
File: SXD_BANG TONG HOP NGHI PHEP NAM 2026.xlsx
Cột B: Họ tên (B9 đến dòng cuối)
Cột C: Ngày vào công chức, cơ quan NN (C9 đến dòng cuối)
"""

import openpyxl
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

# Đọc file Excel
file_path = r'C:\Troly\HTKT\Mau\SXD_bANG TONG HOP NGHI PHEP NAM 2026.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb.active

# Tìm dòng cuối có dữ liệu
last_row = 0
for row in range(9, ws.max_row + 1):
    if ws.cell(row, 2).value:  # Nếu cột B (họ tên) có dữ liệu
        if isinstance(ws.cell(row, 2).value, str) and not ws.cell(row, 2).value.strip().upper().startswith(('I ', 'II ', 'III ')):
            last_row = row
        elif isinstance(ws.cell(row, 2).value, str) and ws.cell(row, 2).value.strip() and not ws.cell(row, 2).value.strip().isupper():
            last_row = row

# Nếu phương pháp trên không hoạt động, dùng cách khác
if last_row == 0:
    for row in range(9, ws.max_row + 1):
        if ws.cell(row, 2).value and ws.cell(row, 3).value:
            last_row = row

print(f"=== DỮ LIỆU TỪ EXCEL (B9:B{last_row}) ===\n")

# Thu thập dữ liệu
data = []
for row in range(9, last_row + 1):
    stt = ws.cell(row, 1).value
    name = ws.cell(row, 2).value
    ngay_vao = ws.cell(row, 3).value

    # Bỏ qua dòng tiêu đề (I, II, III)
    if not name or not isinstance(name, str):
        continue

    name_upper = name.strip().upper()
    if name_upper.startswith(('I ', 'II ', 'III ')) and not any(char.isdigit() for char in name[:5]):
        continue

    data.append({
        'stt': int(stt) if isinstance(stt, (int, float)) else stt,
        'hoTen': name.strip() if isinstance(name, str) else str(name),
        'ngayVao': str(ngay_vao).strip() if ngay_vao else ''
    })

    stt_str = str(int(stt)) if isinstance(stt, (int, float)) else str(stt)
    print(f"{stt_str:3s} | {name:35s} | {ngay_vao}")

print(f"\n=== TẠO CÂU INSERT SQL ===\n")

# Tạo câu INSERT
insert_sql = "INSERT INTO temp_ngay_vao_cong_chuc VALUES\n"
for i, row_data in enumerate(data):
    stt = row_data['stt']
    ho_ten = row_data['hoTen'].replace("'", "''")  # Escape quotes
    ngay_vao = row_data['ngayVao']

    comma = "," if i < len(data) - 1 else ";"
    insert_sql += f"  ({stt}, '{ho_ten}', '{ngay_vao}'){comma}\n"

print(insert_sql)

print(f"\n=== THỐNG KÊ ===")
print(f"Tổng số người: {len(data)}")
print(f"Người có ngày vào: {sum(1 for d in data if d['ngayVao'])}")
print(f"Người không có ngày vào: {sum(1 for d in data if not d['ngayVao'])}")

# Cảnh báo
print("\n⚠ LƯU Ý:")
print("1. Kiểm tra kỹ số liệu họ tên - nếu có lỗi chính tả sẽ không match")
print("2. Nếu có dòng không khớp, KHÔNG TỰ ĐẶT GIÁ TRỊ - HỎI NGƯỜI DÙNG")
print("3. Copy câu INSERT vào file import-ngay-vao-cong-chuc.sql rồi chạy script")

wb.close()
